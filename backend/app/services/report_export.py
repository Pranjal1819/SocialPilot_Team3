# app/services/report_export.py

from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

import openpyxl
from openpyxl.styles import Font, PatternFill

# ==========================================================
# PDF EXPORT
# ==========================================================


def export_report_to_pdf(report_data: dict, report_name: str, file_path: str):

    doc = SimpleDocTemplate(file_path, pagesize=A4)

    styles = getSampleStyleSheet()

    cell_style = ParagraphStyle(
        "CellText", parent=styles["Normal"], fontSize=7, leading=9
    )

    header_style = ParagraphStyle(
        "HeaderText",
        parent=styles["Normal"],
        fontSize=7,
        leading=9,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )

    elements = []

    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"], fontSize=18, spaceAfter=6
    )

    elements.append(Paragraph(report_name, title_style))

    elements.append(
        Paragraph(
            f"Generated on {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}",
            styles["Normal"],
        )
    )

    elements.append(Spacer(1, 0.3 * inch))

    # ------------------------------------------
    # Summary section
    # ------------------------------------------

    elements.append(Paragraph("Summary", styles["Heading2"]))

    summary = report_data.get("summary", {})

    summary_rows = [["Metric", "Value"]]

    for key, value in summary.items():

        label = key.replace("_", " ").title()

        summary_rows.append([label, str(value)])

    summary_table = Table(summary_rows, colWidths=[3 * inch, 2.5 * inch])

    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#f5f5f5")],
                ),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )

    elements.append(summary_table)
    elements.append(Spacer(1, 0.4 * inch))

    # ------------------------------------------
    # Table sections — cell values wrapped in
    # Paragraph so long text wraps instead of
    # overflowing into neighboring columns
    # ------------------------------------------

    tables = report_data.get("tables", {})

    for table_name, rows in tables.items():

        if not rows:
            continue

        section_title = table_name.replace("_", " ").title()

        elements.append(Paragraph(section_title, styles["Heading2"]))
        elements.append(Spacer(1, 0.1 * inch))

        headers = list(rows[0].keys())

        table_data = [
            [Paragraph(h.replace("_", " ").title(), header_style) for h in headers]
        ]

        for row in rows:
            table_data.append(
                [Paragraph(str(row.get(h, "")), cell_style) for h in headers]
            )

        col_width = min(1.3 * inch, 6.5 * inch / len(headers))

        data_table = Table(
            table_data, colWidths=[col_width] * len(headers), repeatRows=1
        )

        data_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#34495e")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#f5f5f5")],
                    ),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )

        elements.append(data_table)
        elements.append(Spacer(1, 0.3 * inch))

    doc.build(elements)


# ==========================================================
# EXCEL EXPORT
# ==========================================================


def export_report_to_excel(report_data: dict, report_name: str, file_path: str):

    wb = openpyxl.Workbook()

    header_fill = PatternFill(
        start_color="2C3E50", end_color="2C3E50", fill_type="solid"
    )
    header_font = Font(color="FFFFFF", bold=True)

    # ------------------------------------------
    # Summary sheet
    # ------------------------------------------

    summary_sheet = wb.active
    summary_sheet.title = "Summary"

    summary_sheet["A1"] = report_name
    summary_sheet["A1"].font = Font(size=14, bold=True)

    summary_sheet["A2"] = (
        f"Generated on {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    )

    summary_sheet["A4"] = "Metric"
    summary_sheet["B4"] = "Value"

    summary_sheet["A4"].fill = header_fill
    summary_sheet["A4"].font = header_font
    summary_sheet["B4"].fill = header_fill
    summary_sheet["B4"].font = header_font

    row = 5

    for key, value in report_data.get("summary", {}).items():

        summary_sheet[f"A{row}"] = key.replace("_", " ").title()
        summary_sheet[f"B{row}"] = value

        row += 1

    summary_sheet.column_dimensions["A"].width = 30
    summary_sheet.column_dimensions["B"].width = 20

    # ------------------------------------------
    # One sheet per table
    # ------------------------------------------

    for table_name, rows in report_data.get("tables", {}).items():

        if not rows:
            continue

        sheet_title = table_name.replace("_", " ").title()[
            :31
        ]  # Excel sheet name limit

        sheet = wb.create_sheet(title=sheet_title)

        headers = list(rows[0].keys())

        for col_idx, header in enumerate(headers, start=1):

            cell = sheet.cell(
                row=1, column=col_idx, value=header.replace("_", " ").title()
            )
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, row_data in enumerate(rows, start=2):

            for col_idx, header in enumerate(headers, start=1):

                sheet.cell(row=row_idx, column=col_idx, value=row_data.get(header))

        for col_idx in range(1, len(headers) + 1):
            sheet.column_dimensions[chr(64 + col_idx)].width = 20

    wb.save(file_path)
